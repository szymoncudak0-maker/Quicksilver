"""Quicksilver ONLINE 2.0: sesje użytkowników, prywatne zadania i czasowe wyniki.

Wymaga HTTPS oraz hasła dostępu ustawionego w QS_ACCESS_PASSWORD.
Jedna instancja aplikacji / jeden worker. Do pracy na wielu replikach
potrzebne byłoby współdzielone repozytorium plików i kolejka zadań.
"""
from __future__ import annotations

import json
import os
import secrets
import shutil
import tempfile
import threading
import time
import uuid
from collections import defaultdict
from dataclasses import asdict
from pathlib import Path
from zipfile import ZipFile, ZIP_STORED

from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from openpyxl.utils import get_column_letter
from pydantic import BaseModel
from starlette.staticfiles import StaticFiles

from engine import Settings, load_source, run
from safe_images import download_public_image
from settlements import create_router

ROOT = Path(__file__).resolve().parent
WORK = Path(tempfile.mkdtemp(prefix="quicksilver_online_"))
MAX_UPLOAD = 25 * 1024 * 1024
SESSION_LIFETIME = 12 * 3600
JOB_LIFETIME = 2 * 3600
MAX_JOBS = max(1, min(8, int(os.getenv("QS_MAX_JOBS", "2"))))
SECURE_COOKIE = os.getenv("QS_SECURE_COOKIES", "true").lower() == "true"
PASSWORD = os.getenv("QS_ACCESS_PASSWORD", "")
COOKIE_NAME = "quicksilver_sid"
LOCK = threading.RLock()
SEMAPHORE = threading.BoundedSemaphore(MAX_JOBS)
SESSIONS: dict[str, dict] = {}
UPLOADS: dict[str, dict] = {}
JOBS: dict[str, dict] = {}
LOGIN_ATTEMPTS: dict[str, list[float]] = defaultdict(list)
app = FastAPI(title="Quicksilver ONLINE", docs_url=None, redoc_url=None, openapi_url=None)


@app.middleware("http")
async def standard_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self'; img-src 'self' data:; style-src 'self'; script-src 'self'; "
        "connect-src 'self'; base-uri 'none'; form-action 'self'; frame-ancestors 'none'")
    return response


def fail(status, message):
    raise HTTPException(status_code=status, detail=message)


def tidy():
    """Usuń stare sesje, przesłane pliki i wyniki, ale nie prace w toku."""
    now = time.time()
    paths = []
    with LOCK:
        for sid, sess in list(SESSIONS.items()):
            if now - sess["created"] > SESSION_LIFETIME:
                del SESSIONS[sid]
        for key, upload in list(UPLOADS.items()):
            if now - upload["created"] > JOB_LIFETIME:
                paths.append(upload["file"])
                del UPLOADS[key]
        for key, job in list(JOBS.items()):
            if job.get("finished") and now - job["finished"] > JOB_LIFETIME:
                paths.append(job["work"])
                del JOBS[key]
        for key, attempts in list(LOGIN_ATTEMPTS.items()):
            fresh = [t for t in attempts if now - t < 300]
            if fresh:
                LOGIN_ATTEMPTS[key] = fresh
            else:
                del LOGIN_ATTEMPTS[key]
    for p in paths:
        if p.is_dir():
            shutil.rmtree(p, ignore_errors=True)
        else:
            p.unlink(missing_ok=True)


def need_session(request: Request):
    sid = request.cookies.get(COOKIE_NAME, "")
    with LOCK:
        sess = SESSIONS.get(sid)
        if not sess or time.time() - sess["created"] > SESSION_LIFETIME:
            fail(401, "Sesja wygasła. Zaloguj się ponownie.")
        return sid, sess


def need_write(request: Request, x_csrf_token: str | None = Header(default=None)):
    sid, sess = need_session(request)
    if not x_csrf_token or not secrets.compare_digest(x_csrf_token, sess["csrf"]):
        fail(403, "Nieprawidłowa ochrona sesji. Odśwież stronę.")
    origin = request.headers.get("origin")
    if origin:
        expected = os.getenv("QS_PUBLIC_ORIGIN", "").strip().rstrip("/") or str(request.base_url).rstrip("/")
        if origin.rstrip("/") != expected:
            fail(403, "Niedozwolone pochodzenie żądania.")
    return sid, sess


def get_upload(upload_id, sid):
    with LOCK:
        upload = UPLOADS.get(str(upload_id))
        if not upload or upload["sid"] != sid:
            fail(404, "Nie znaleziono przesłanego pliku.")
        return upload["file"]


def get_job(job_id, sid):
    with LOCK:
        job = JOBS.get(job_id)
        if not job or job["sid"] != sid:
            fail(404, "Nie znaleziono zadania.")
        return job


def cell_view(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target[:180]
    return str(cell.value)[:180] if cell.value is not None else ""


def detect_photo_columns(ws):
    """Suggest EAN and photo columns from column names without overwriting data."""
    headers = [str(ws.cell(1, col).value or '').lower().replace('_', ' ')
               for col in range(1, min(ws.max_column or 0, 60) + 1)]
    ean_idx = next((i + 1 for i, name in enumerate(headers)
                    if any(word in name.split() for word in ('ean', 'gtin'))), None)
    photos = [i + 1 for i, name in enumerate(headers)
              if any(word in name for word in ('zdjec', 'zdjęc', 'photo', 'image', 'foto'))]
    if not ean_idx or not photos:
        return None
    first = next((i for i in photos if i != ean_idx), None)
    if not first:
        return None
    photo_set = set(photos)
    columns = 0
    while columns < 11 and (first + columns) in photo_set:
        columns += 1
    return {'ean_column': get_column_letter(ean_idx),
            'photo_column': get_column_letter(first), 'photo_columns': columns}


def preview(path: Path, sheet: str | None = None, start=1):
    wb = load_source(path, Settings(verify_tls=True))
    try:
        sheets = list(wb.sheetnames)
        if not sheets:
            raise ValueError("Brak arkuszy Excela.")
        if sheet not in sheets:
            sheet = sheets[0]
        ws = wb[sheet]
        first = max(1, min(1_048_576, int(start)))
        cols = min(ws.max_column or 2, 14)
        rows = [[cell_view(c) for c in row] for row in ws.iter_rows(
            min_row=first, max_row=min(ws.max_row, first + 5), max_col=cols)]
        return {"sheets": sheets, "sheet": sheet, "rows": rows, "start_row": first,
                "total_rows": ws.max_row, "visible_cols": cols,
                "detected_columns": detect_photo_columns(ws)}
    finally:
        wb.close()


def read_settings(data: dict):
    if not isinstance(data, dict):
        raise ValueError("Nieprawidłowe ustawienia.")
    fields = {"ean_column": str, "photo_column": str, "photo_columns": int, "max_photos": int,
              "first_row": int, "separator": str, "output_format": str, "timeout": int,
              "retries": int, "split_batches": bool, "batch_limit_mb": float,
              "package_type": str, "csv_delimiter": str, "csv_encoding": str}
    values = {}
    for field, cast in fields.items():
        if field in data:
            if field == "split_batches" and type(data[field]) is not bool:
                raise ValueError("Nieprawidłowy parametr podziału paczek.")
            values[field] = cast(data[field])
    values["ean_column"] = values.get("ean_column", "A").strip().upper()
    values["photo_column"] = values.get("photo_column", "B").strip().upper()
    values["verify_tls"] = True  # W chmurze nigdy nie omijamy certyfikatów HTTPS.
    values["layout"] = "product"
    cfg = Settings(**{**asdict(Settings()), **values})
    cfg.validate()
    return cfg


def notify(job, event, *args):
    with LOCK:
        if event == "log":
            # log zawiera wybrane przez użytkownika nazwy obrazów / URL
            job["logs"].append(str(args[0])[:400])
            job["logs"] = job["logs"][-140:]
        elif event == "progress":
            n, total, photos, errors = args
            job.update({"processed": n, "total": total, "photos": photos, "errors": errors,
                        "progress": round(100 * n / max(total, 1))})


def zip_files(folder, target):
    with ZipFile(target, "w", compression=ZIP_STORED, allowZip64=True) as zipf:
        for item in sorted(folder.rglob("*")):
            if item.is_file():
                zipf.write(item, arcname=item.relative_to(folder).as_posix())


def worker(job_id, source, sheet, settings):
    with LOCK:
        job = JOBS[job_id]
        job["state"] = "running"
    try:
        result = run(source, sheet, job["work"], settings,
                     emit=lambda *args: notify(job, *args), stop=job["stop"],
                     fetch=download_public_image)
        with LOCK:
            job.update(result)
            job["state"] = "packing"
            job["logs"].append("Tworzenie archiwum do pobrania…")
        folder = Path(result["folder"])
        package_path = job["work"] / f"{folder.name}_wyniki.zip"
        zip_files(folder, package_path)
        assets = []
        for f in sorted(folder.rglob("wynik.xlsx")) + sorted(folder.rglob("zdjecia.zip")):
            assets.append({"path": f, "filename": f.name, "location": f.parent.relative_to(folder).as_posix(),
                           "kind": "excel" if f.name == "wynik.xlsx" else "photos"})
        with LOCK:
            job["bundle"] = package_path
            job["artifacts"] = [{k:v for k,v in a.items() if k != "path"} | {"index": i}
                                for i,a in enumerate(assets)]
            job["files"] = [a["path"] for a in assets]
            job["state"] = "cancelled" if result["cancelled"] else "complete"
            job["progress"] = 100
    except Exception as exc:
        with LOCK:
            job["state"] = "failed"
            detail = str(exc).strip() or type(exc).__name__
            job["message"] = detail[:300]
            job["logs"].append("Przetwarzanie nie powiodło się: " + job["message"])
            job["logs"].append("Sprawdź wskazane kolumny, dostępność zdjęć i limit 500 MB na zadanie.")
    finally:
        with LOCK:
            job["finished"] = time.time()
        SEMAPHORE.release()


def job_status(job):
    allowed = ("id", "state", "progress", "processed", "total", "photos", "products", "errors",
               "batches", "oversized_batches", "cancelled", "message", "logs", "artifacts")
    result = {key: job[key] for key in allowed if key in job}
    if job.get("folder"):
        result["folder"] = Path(job["folder"]).name
    return result


class Login(BaseModel):
    password: str


class SheetPreview(BaseModel):
    upload_id: str
    sheet: str | None = None
    start: int = 1


class Start(BaseModel):
    upload_id: str
    sheet: str
    settings: dict = {}


@app.get("/health")
def health():
    return {"ok": True}


@app.post("/api/login")
def login(data: Login, request: Request):
    tidy()
    if not PASSWORD or len(PASSWORD) < 12:
        fail(503, "Administrator musi ustawić bezpieczne hasło QS_ACCESS_PASSWORD (min. 12 znaków).")
    remote = request.client.host if request.client else "?"
    with LOCK:
        attempt = LOGIN_ATTEMPTS[remote]
        if len(attempt) >= 10:
            fail(429, "Zbyt wiele prób logowania. Spróbuj za 5 minut.")
        if not secrets.compare_digest(data.password.encode(), PASSWORD.encode()):
            attempt.append(time.time())
            fail(401, "Nieprawidłowe hasło.")
        attempt.clear()
        sid = secrets.token_urlsafe(36)
        SESSIONS[sid] = {"created": time.time(), "csrf": secrets.token_urlsafe(32)}
    response = JSONResponse({"ok": True})
    response.set_cookie(COOKIE_NAME, sid, max_age=SESSION_LIFETIME, path="/", secure=SECURE_COOKIE,
                        httponly=True, samesite="strict")
    return response


@app.post("/api/logout")
def logout(request: Request, session=Depends(need_write)):
    sid, _ = session
    with LOCK:
        SESSIONS.pop(sid, None)
    response = JSONResponse({"ok": True})
    response.delete_cookie(COOKIE_NAME, path="/", secure=SECURE_COOKIE, httponly=True, samesite="strict")
    return response


@app.get("/api/bootstrap")
def bootstrap(session=Depends(need_session)):
    _, sess = session
    return {"key": sess["csrf"], "defaults": asdict(Settings(verify_tls=True)),
            "version": "2.0", "output_base": "Chmura · wyniki dostępne do pobrania przez 2 godziny"}


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...), session=Depends(need_write)):
    sid, _ = session
    tidy()
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".xlsx", ".xlsm", ".csv"):
        fail(400, "Dozwolone formaty: XLSX, XLSM, CSV.")
    identity = uuid.uuid4().hex
    target = WORK / (identity + suffix)
    size = 0
    try:
        with target.open("xb") as out:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_UPLOAD:
                    fail(413, "Plik przekracza limit 25 MB.")
                out.write(chunk)
        if not size:
            fail(400, "Nie przesłano zawartości pliku.")
        data = preview(target)
    except HTTPException:
        target.unlink(missing_ok=True)
        raise
    except Exception as exc:
        target.unlink(missing_ok=True)
        fail(400, f"Nie można odczytać pliku: {str(exc)[:200]}")
    finally:
        await file.close()
    with LOCK:
        UPLOADS[identity] = {"sid": sid, "file": target, "created": time.time()}
    return {"upload_id": identity, "file_name": Path(file.filename).name, **data}


@app.post("/api/preview")
def preview_route(data: SheetPreview, session=Depends(need_write)):
    sid, _ = session
    source = get_upload(data.upload_id, sid)
    try:
        return preview(source, data.sheet, data.start)
    except Exception as exc:
        fail(400, f"Nie można wyświetlić podglądu: {str(exc)[:200]}")


@app.post("/api/start", status_code=202)
def start_route(data: Start, session=Depends(need_write)):
    sid, _ = session
    source = get_upload(data.upload_id, sid)
    try:
        settings = read_settings(data.settings)
        if data.sheet not in preview(source)["sheets"]:
            fail(400, "Wybierz arkusz istniejący w pliku.")
    except (TypeError, ValueError) as exc:
        fail(400, str(exc))
    if not SEMAPHORE.acquire(blocking=False):
        fail(429, "Serwer przetwarza maksymalną liczbę zadań. Spróbuj za chwilę.")
    job_id = uuid.uuid4().hex
    work = WORK / job_id
    try:
        work.mkdir(mode=0o700)
        job = {"id": job_id, "sid": sid, "work": work, "created": time.time(),
               "state": "waiting", "progress": 0, "processed": 0, "total": 0, "photos": 0,
               "errors": 0, "products": 0, "logs": ["Rozpoczynanie zadania…"],
               "stop": threading.Event(), "message": ""}
        with LOCK:
            JOBS[job_id] = job
        t = threading.Thread(target=worker, args=(job_id, source, data.sheet, settings), daemon=True)
        t.start()
    except Exception:
        SEMAPHORE.release()
        shutil.rmtree(work, ignore_errors=True)
        raise
    return {"job_id": job_id}


@app.get("/api/jobs/{job_id}")
def status_route(job_id: str, session=Depends(need_session)):
    sid, _ = session
    return job_status(get_job(job_id, sid))


@app.post("/api/jobs/{job_id}/cancel")
def cancel_route(job_id: str, session=Depends(need_write)):
    sid, _ = session
    job = get_job(job_id, sid)
    with LOCK:
        if job["state"] in ("waiting", "running"):
            job["stop"].set()
            job["logs"].append("Zlecono zatrzymanie pobierania.")
    return {"ok": True}


@app.get("/api/jobs/{job_id}/download")
def bundle_route(job_id: str, session=Depends(need_session)):
    sid, _ = session
    job = get_job(job_id, sid)
    if job.get("state") not in ("complete", "cancelled") or not job.get("bundle", Path("/")).is_file():
        fail(404, "Paczka wynikowa nie jest jeszcze gotowa.")
    return FileResponse(job["bundle"], media_type="application/zip", filename=job["bundle"].name)


@app.get("/api/jobs/{job_id}/assets/{index}")
def asset_route(job_id: str, index: int, session=Depends(need_session)):
    sid, _ = session
    job = get_job(job_id, sid)
    files = job.get("files", [])
    if job.get("state") not in ("complete", "cancelled") or not 0 <= index < len(files):
        fail(404, "Nie znaleziono pliku wynikowego.")
    file = files[index]
    if not file.is_file():
        fail(404, "Plik nie jest dostępny.")
    return FileResponse(file, filename=file.name)


app.include_router(create_router(need_session, need_write))

app.mount("/", StaticFiles(directory=ROOT / "static", html=True), name="static")
