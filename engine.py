"""Silnik Quicksilver. Nie wymaga interfejsu graficznego."""
from __future__ import annotations

import io
import os
import csv
import ssl
import re
import time
import uuid
import warnings
import math
from dataclasses import dataclass
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from threading import Event
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageOps


@dataclass
class Settings:
    ean_column: str = "A"
    photo_column: str = "B"
    photo_columns: int = 11
    max_photos: int = 11
    first_row: int = 2
    separator: str = "; "
    layout: str = "product"  # Odczyt starych ustawień; raport zawsze ma jeden wiersz na EAN.
    output_format: str = "webp_png"
    image_policy_version: int = 1
    timeout: int = 20
    retries: int = 2
    verify_tls: bool = False
    csv_delimiter: str = "auto"
    csv_encoding: str = "auto"
    split_batches: bool = True
    batch_limit_mb: float = 150.0
    package_type: str = "zip"

    def validate(self):
        if self.package_type not in ("zip", "folder"):
            raise ValueError("Format paczki musi być ZIP lub folder.")
        if not isinstance(self.split_batches, bool):
            raise ValueError("Nieprawidłowe ustawienie podziału na paczki.")
        if not math.isfinite(self.batch_limit_mb) or not 0.001 <= self.batch_limit_mb <= 100000:
            raise ValueError("Limit paczki musi wynosić od 0,001 do 100000 MB.")
        if self.csv_delimiter not in ("auto", ";", ",", "\t", "|"):
            raise ValueError("Nieprawidłowy separator CSV.")
        if self.csv_encoding not in ("auto", "utf-8-sig", "cp1250", "utf-16", "iso-8859-2"):
            raise ValueError("Nieprawidłowe kodowanie CSV.")
        for col in (self.ean_column, self.photo_column):
            if not re.fullmatch(r"[A-Z]{1,3}", col):
                raise ValueError("Kolumny wpisz literami, np. A, B lub AA.")
            if column_index_from_string(col) > 16384:
                raise ValueError("Kolumna wykracza poza zakres Excela.")
        if not 1 <= self.max_photos <= 100:
            raise ValueError("Limit zdjęć musi wynosić od 1 do 100.")
        if not 1 <= self.photo_columns <= 500:
            raise ValueError("Liczba kolumn zdjęć musi wynosić od 1 do 500.")
        start = column_index_from_string(self.photo_column)
        if start + self.photo_columns - 1 > 16384:
            raise ValueError("Zakres zdjęć wykracza poza arkusz.")
        if start <= column_index_from_string(self.ean_column) < start + self.photo_columns:
            raise ValueError("Kolumna EAN nie może należeć do zakresu zdjęć.")
        if not 1 <= self.first_row <= 1048576:
            raise ValueError("Nieprawidłowy pierwszy wiersz danych.")
        if not 5 <= self.timeout <= 120 or not 0 <= self.retries <= 5:
            raise ValueError("Limit czasu: 5–120 s; ponowienia: 0–5.")
        if self.layout not in ("product", "photo") or self.output_format not in ("jpg", "original", "webp_png", "webp_jpg"):
            raise ValueError("Nieprawidłowy format wyjściowy.")
        if not self.separator or len(self.separator) > 10:
            raise ValueError("Separator musi mieć 1–10 znaków.")


def load_source(path, settings=None, read_only=False):
    """Wspólny odczyt Excela i CSV. CSV zachowuje identyfikatory jako tekst."""
    settings = settings or Settings()
    if Path(path).suffix.lower() != ".csv":
        wb = load_workbook(path, read_only=read_only, data_only=False, keep_links=False)
        # Eksporty XLSX z niektórych systemów nie zapisują elementu <dimension>.
        # Odczyt strumieniowy zwraca wtedy None zamiast liczby wierszy/kolumn.
        # Zwykły odczyt ustala zakres z komórek i obsługuje również pusty arkusz.
        if read_only and any(ws.max_row is None or ws.max_column is None for ws in wb.worksheets):
            wb.close()
            return load_workbook(path, read_only=False, data_only=False, keep_links=False)
        return wb
    raw = Path(path).read_bytes()
    if not raw.strip():
        raise ValueError("Plik CSV jest pusty.")
    encoding = settings.csv_encoding
    if encoding == "auto":
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encoding = "utf-16"
        else:
            try:
                raw.decode("utf-8-sig")
                encoding = "utf-8-sig"
            except UnicodeDecodeError:
                encoding = "cp1250"
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError as exc:
        raise ValueError("Nie można odczytać CSV. Wybierz kodowanie w ustawieniach.") from exc
    if "\x00" in text:
        raise ValueError("Nieprawidłowe kodowanie CSV. Wybierz UTF-16 w ustawieniach.")
    delimiter = settings.csv_delimiter
    if delimiter == "auto":
        try:
            delimiter = csv.Sniffer().sniff(text[:131072], delimiters=";,\t|").delimiter
        except csv.Error:
            # Przy nieregularnych wierszach preferuj separator z nagłówka.
            first = text.splitlines()[0]
            delimiter = max(";,\t|", key=first.count) if any(c in first for c in ";,\t|") else ";"
    csv.field_size_limit(10 * 1024 * 1024)
    wb = Workbook()
    ws = wb.active
    ws.title = "CSV"
    try:
        for row_index, values in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True), 1):
            if row_index > 1048576 or len(values) > 16384:
                raise ValueError("CSV przekracza zakres arkusza Excel.")
            for col_index, value in enumerate(values, 1):
                cell = ws.cell(row_index, col_index, value)
                cell.data_type = "s"
        wb.flash_csv_info = f"CSV · {encoding} · separator: {repr(delimiter)}"
        return wb
    except Exception:
        wb.close()
        raise


def ean_text(cell):
    value = cell.value
    if value is None or str(value).strip() == "":
        return ""
    if isinstance(value, bool):
        raise ValueError("EAN nie jest kodem liczbowym.")
    if isinstance(value, (int, float)):
        if int(value) != value or value < 0:
            raise ValueError("EAN musi składać się z cyfr.")
        result = str(int(value))
        if re.fullmatch(r"0+", cell.number_format or ""):
            result = result.zfill(len(cell.number_format))
    else:
        result = str(value).strip()
    if not re.fullmatch(r"[0-9]{8,14}", result):
        raise ValueError("EAN/GTIN musi mieć 8–14 cyfr. Zapisz kod w Excelu jako tekst.")
    return result


def cell_urls(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return [cell.hyperlink.target.strip()]
    value = str(cell.value or "").strip()
    if not value:
        return []
    # HYPERLINK z adresem literalnym; pozostałe formuły wymagają przeliczenia w Excelu.
    if value.startswith("="):
        match = re.match(r'=HYPERLINK\(\s*"([^"]+)"', value, flags=re.I)
        if not match:
            raise ValueError("Formuła linku nie jest obsługiwana. Wklej adres jako wartość.")
        value = match.group(1)
    # Średnik/przecinek rozdziela adresy tylko wtedy, gdy zaczyna się kolejny URL.
    return [x.strip() for x in re.split(r"[\r\n]+|[;,|]\s*(?=https?://)", value) if x.strip()]


class Cancelled(Exception):
    pass


def download(url, settings, stop):
    parts = urlsplit(url)
    if parts.scheme not in ("http", "https") or not parts.hostname:
        raise ValueError("Wymagany bezpośredni adres zdjęcia HTTP lub HTTPS.")
    limit = 40 * 1024 * 1024
    for attempt in range(settings.retries + 1):
        if stop.is_set():
            raise Cancelled()
        try:
            request = Request(url, headers={"User-Agent": "Mozilla/5.0 (compatible; Quicksilver/3.0)", "Accept": "image/*,*/*;q=0.8"})
            started = time.monotonic()
            context = ssl.create_default_context()
            if not settings.verify_tls:
                context.check_hostname = False
                context.verify_mode = ssl.CERT_NONE
            with urlopen(request, timeout=settings.timeout, context=context) as response:
                if urlsplit(response.url).scheme not in ("http", "https"):
                    raise ValueError("Niedozwolone przekierowanie.")
                if int(response.headers.get("Content-Length", "0")) > limit:
                    raise ValueError("Zdjęcie przekracza limit 40 MB.")
                data = bytearray()
                while True:
                    if stop.is_set():
                        raise Cancelled()
                    if time.monotonic() - started > settings.timeout:
                        raise TimeoutError("Przekroczono czas pobierania zdjęcia.")
                    chunk = response.read(64 * 1024)
                    if not chunk:
                        break
                    data.extend(chunk)
                    if len(data) > limit:
                        raise ValueError("Zdjęcie przekracza limit 40 MB.")
            return bytes(data)
        except (URLError, TimeoutError, OSError) as exc:
            if isinstance(exc, HTTPError) and exc.code not in (408, 429, 500, 502, 503, 504):
                raise
            if attempt == settings.retries:
                raise
            if stop.wait(min(2 ** attempt, 4)):
                raise Cancelled()


def image_bytes(raw, mode):
    extensions = {"JPEG": ".jpg", "PNG": ".png", "WEBP": ".webp", "GIF": ".gif", "BMP": ".bmp", "TIFF": ".tif"}
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(raw)) as check:
            fmt = check.format
            check.verify()
        if fmt not in extensions:
            raise ValueError("Nieobsługiwany format zdjęcia (np. SVG/AVIF).")
        with Image.open(io.BytesIO(raw)) as source:
            source.load()
            if mode == "original" or (mode in ("webp_png", "webp_jpg") and fmt != "WEBP"):
                return raw, extensions[fmt]
            fixed = ImageOps.exif_transpose(source).convert("RGBA")
            if mode == "webp_png":
                buf = io.BytesIO()
                fixed.save(buf, "PNG")
                return buf.getvalue(), ".png"
            background = Image.new("RGB", fixed.size, "white")
            background.paste(fixed, mask=fixed.getchannel("A"))
            buf = io.BytesIO()
            background.save(buf, "JPEG", quality=95, optimize=True)
            return buf.getvalue(), ".jpg"


def export_report(folder, products, errors, settings):
    wb = Workbook()
    ws = wb.active
    ws.title = "Zdjęcia"
    ws.append(["Typ", "EAN", "Zdjęcie okładki/produktu", "Dodatkowe zdjęcia"])
    for ean, names in products.items():
        ws.append(["singiel", ean, names[0] if names else "", settings.separator.join(names[1:])])
    issues = wb.create_sheet("Błędy")
    issues.append(["Wiersz źródłowy", "EAN", "Kolumna / link", "Powód"])
    for entry in errors:
        issues.append(list(entry))
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="F58220")
            cell.font = Font(bold=True, color="161616")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Nie wykonuj formuł pochodzących z danych.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions["A"].width = 23
        sheet.column_dimensions["B"].width = 23
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 90
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = max(20, 15 * math.ceil(len(str(ws.cell(row, 4).value or "")) / 85))
    issues.column_dimensions["C"].width = 60
    issues.column_dimensions["D"].width = 65
    for cell in ws["B"][1:]:
        cell.number_format = "@"
    temporary = folder / "wynik.tmp.xlsx"
    wb.save(temporary)
    temporary.replace(folder / "wynik.xlsx")
    wb.close()


def plan_batches(products, photos, limit_bytes):
    """EAN jest niepodzielną grupą, także gdy występował w odległych wierszach."""
    batches, current, size = [], {}, 0
    for ean, names in products.items():
        group_size = sum((photos / name).stat().st_size for name in names)
        if current and size + group_size > limit_bytes:
            batches.append((current, size))
            current, size = {}, 0
        current[ean] = names
        size += group_size
        if group_size > limit_bytes:
            batches.append((current, size))
            current, size = {}, 0
    if current:
        batches.append((current, size))
    return batches or [({}, 0)]


def archive_photos(folder, names):
    """Usuń robocze obrazy dopiero po zapisaniu i sprawdzeniu kompletnego ZIP-a."""
    photos = folder / "zdjecia"
    names = list(names)
    temporary = folder / "zdjecia.tmp.zip"
    destination = folder / "zdjecia.zip"
    try:
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED, compresslevel=6) as archive:
            for name in names:
                archive.write(photos / name, arcname=name)
        with ZipFile(temporary) as archive:
            if archive.namelist() != names or archive.testzip() is not None:
                raise OSError("Nie udało się potwierdzić kompletności archiwum zdjęć.")
            for info in archive.infolist():
                if info.file_size != (photos / info.filename).stat().st_size:
                    raise OSError("Rozmiar zdjęcia w ZIP-ie nie zgadza się z plikiem źródłowym.")
        temporary.replace(destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    for name in names:
        (photos / name).unlink()
    photos.rmdir()
    return destination


def split_output(folder, products, errors, settings, emit):
    photos = folder / "zdjecia"
    limit_bytes = round(settings.batch_limit_mb * 1_000_000)
    batches = plan_batches(products, photos, limit_bytes)
    oversized = 0
    notes = [f"Limit zdjęć: {settings.batch_limit_mb:g} MB (1 MB = 1 000 000 bajtów).",
             "Rozmiar Exceli nie jest wliczany. Liczone są zapisane zdjęcia po konwersji.",
             ("Zdjęcia w każdej paczce znajdują się w zdjecia.zip." if settings.package_type == "zip" else "Zdjęcia w każdej paczce znajdują się w folderze zdjecia."),
             "Limit dotyczy sumy wielkości plików zdjęć, a nie skompresowanego ZIP-a.",
             "Wszystkie pobrane zdjęcia danego EAN-u należą do jednej paczki.", ""]
    # Zapis planu przed przenoszeniem ułatwia odnalezienie plików przy błędzie dysku.
    import json
    plan = [{"folder": f"paczka_{i:03d}", "bytes": size, "products": group}
            for i, (group, size) in enumerate(batches, 1)]
    (folder / "podzial.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    for index, (group, size) in enumerate(batches, 1):
        batch = folder / f"paczka_{index:03d}"
        batch.mkdir()
        target = batch / "zdjecia"
        target.mkdir()
        batch_errors = [e for e in errors if e[1] in group or (index == 1 and e[1] not in products)]
        if size > limit_bytes:
            oversized += 1
            warning = f"EAN przekracza limit paczki: {size / 1_000_000:.3f} MB. Zdjęcia pozostawiono razem."
            batch_errors.append(("", next(iter(group)), "", warning))
            emit("log", warning)
        export_report(batch, group, batch_errors, settings)
        for names in group.values():
            for name in names:
                (photos / name).replace(target / name)
        if settings.package_type == "zip":
            archive_photos(batch, [name for names in group.values() for name in names])
        line = f"{batch.name}: {size / 1_000_000:.3f} MB | EAN-y: {len(group)} | zdjęcia: {sum(map(len, group.values()))}"
        if size > limit_bytes:
            line += " | PRZEKROCZONY LIMIT (jeden EAN)"
        notes.append(line)
        emit("log", line)
    (folder / "podsumowanie.txt").write_text("\n".join(notes) + "\n", encoding="utf-8")
    photos.rmdir()
    (folder / "wynik.xlsx").unlink()
    return len(batches), oversized


def max_output_bytes():
    """Total photo-data ceiling; 1 MB means 1,000,000 bytes."""
    return max(1, int(os.getenv("QS_MAX_OUTPUT_MB", "5000"))) * 1_000_000


def run(source, sheet_name, output, settings, emit=lambda *a: None, stop=None, fetch=download):
    settings.validate()
    stop = stop or Event()
    wb = load_source(source, settings)
    try:
        ws = wb[sheet_name]
        if settings.first_row > ws.max_row:
            raise ValueError("Pierwszy wiersz danych jest poza arkuszem.")
        folder = Path(output) / (datetime.now(ZoneInfo("Europe/Warsaw")).strftime("Quicksilver_%Y%m%d_%H%M%S_") + uuid.uuid4().hex[:6])
        folder.mkdir(parents=True, exist_ok=False)
        photos = folder / "zdjecia"
        photos.mkdir()
        products, seen, errors = {}, {}, []
        total = ws.max_row - settings.first_row + 1
        ean_col = column_index_from_string(settings.ean_column)
        start = column_index_from_string(settings.photo_column)
        cancelled = False
        count = 0
        result_bytes = 0
        max_result_bytes = max_output_bytes()
        batch_count, oversized = 1, 0
        def issue(row, ean, url, reason):
            errors.append((row, ean, url, reason))
            emit("log", f"Wiersz {row} · {ean or 'brak EAN'} · {reason}")
        try:
            for n, row in enumerate(range(settings.first_row, ws.max_row + 1), 1):
                if stop.is_set():
                    raise Cancelled()
                cell = ws.cell(row, ean_col)
                photo_cells = [ws.cell(row, c) for c in range(start, start + settings.photo_columns)]
                if cell.value is None and not any(c.value or c.hyperlink for c in photo_cells):
                    emit("progress", n, total, count, len(errors))
                    continue
                try:
                    ean = ean_text(cell)
                    if not ean:
                        raise ValueError("Brak EAN.")
                except ValueError as exc:
                    issue(row, str(cell.value or ""), get_column_letter(ean_col), str(exc))
                    emit("progress", n, total, count, len(errors))
                    continue
                names = products.setdefault(ean, [])
                known = seen.setdefault(ean, set())
                had_links = False
                for c in photo_cells:
                    if len(names) >= settings.max_photos:
                        break
                    try:
                        urls = cell_urls(c)
                    except ValueError as exc:
                        issue(row, ean, c.column_letter, str(exc))
                        continue
                    had_links |= bool(urls)
                    for url in urls:
                        if stop.is_set():
                            raise Cancelled()
                        if len(names) >= settings.max_photos:
                            break
                        if url in known:
                            continue
                        try:
                            raw, extension = image_bytes(fetch(url, settings, stop), settings.output_format)
                        except Cancelled:
                            raise
                        except Exception as exc:
                            issue(row, ean, url, f"Nie pobrano zdjęcia: {exc}")
                            continue
                        if result_bytes + len(raw) > max_result_bytes:
                            raise RuntimeError(f"Przekroczono limit {max_result_bytes / 1_000_000:g} MB zdjęć dla jednego zadania. Podziel plik źródłowy na mniejsze części.")
                        result_bytes += len(raw)
                        index = len(names) + 1
                        name = ean + (f"_{index}" if index > 1 else "") + extension
                        target = photos / name
                        # Błędy dysku przerywają proces; nie są błędami linków.
                        with target.open("xb") as handle:
                            handle.write(raw)
                        names.append(name)
                        known.add(url)
                        count += 1
                        emit("log", f"Zapisano {name}")
                if not had_links and not names:
                    issue(row, ean, "", "Brak linków w wybranym zakresie kolumn.")
                emit("progress", n, total, count, len(errors))
        except Cancelled:
            cancelled = True
        finally:
            export_report(folder, products, errors, settings)
        if settings.split_batches:
            emit("log", "Tworzenie paczek i Exceli — całe EAN-y pozostają razem…")
            batch_count, oversized = split_output(folder, products, errors, settings, emit)
        else:
            if settings.package_type == "zip":
                emit("log", "Pakowanie zdjęć do zdjecia.zip…")
                archive_photos(folder, [name for names in products.values() for name in names])
        # Folder jest nazywany dopiero po skończeniu całej pracy, nie w momencie startu.
        finished = datetime.now(ZoneInfo("Europe/Warsaw")).strftime("Quicksilver_%Y-%m-%d_%H-%M-%S")
        final_folder = Path(output) / finished
        suffix = 2
        while final_folder.exists():
            final_folder = Path(output) / f"{finished}_{suffix}"
            suffix += 1
        folder.rename(final_folder)
        folder = final_folder
        emit("log", f"Zakończono. Wyniki: {folder}")
        return {"folder": str(folder), "photos": count, "products": len(products), "errors": len(errors),
                "cancelled": cancelled, "batches": batch_count, "oversized_batches": oversized}
    finally:
        wb.close()
